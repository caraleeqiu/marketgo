import openpyxl, warnings, json, datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
warnings.filterwarnings("ignore")

SRC="/root/.claude/uploads/38575678-d9c2-5e63-b431-13963f45edc5/77c23b2e-iLands_Intentional_Shaper_TikTok_KOC_300_with_Emails0620.xlsx"
OUT="/home/user/marketgo/iLands_Intentional_Shaper_TikTok_KOC_350_with_Emails_2026-06-22.xlsx"

# dedup check
existing=set(json.load(open("/home/user/marketgo/.research/existing_handles.json")))

# ---- cluster-level templates ----
SHAPE_LINE={
 "Reflection & Writing":"You already turn inner life into a practice.",
 "Worldbuilding & Story":"You already author whole worlds and make them feel inhabited.",
 "Visual & Virtual Identity":"You already design and inhabit a deliberate visual identity.",
 "Builder & Creative AI":"You already build things in public and shape them over time.",
 "Cozy & Systems":"You already turn intentional routines into living systems.",
}
ARC={
 "Reflection & Writing":"14-night shaping diary: one 20-minute evening reflection becomes a visible choice, memory or relationship change in the Agent.",
 "Worldbuilding & Story":"14-day worldbuilding diary: each session adds canon, characters and history that your AI resident remembers and acts on.",
 "Visual & Virtual Identity":"14-day identity diary: each session shapes how your AI resident looks, presents and evolves its persona.",
 "Builder & Creative AI":"14-day build-in-public diary: each session ships a visible change to an AI resident you are raising from scratch.",
 "Cozy & Systems":"14-day cozy-systems diary: one calm evening ritual becomes a routine your AI resident adopts, remembers and refines.",
}
SHAPER_VALUE={
 "Reflection & Writing":"验证 iLands 是否能成为一种会复利的数字关系：Parent 把判断、人生语境与反思沉淀进一个持续成长的 AI。",
 "Worldbuilding & Story":"验证 Parent 是否会把世界观、设定与叙事记忆投入一个持续演化、能记住 canon 的 AI 居民。",
 "Visual & Virtual Identity":"验证 Parent 是否愿意把虚拟形象与身份审美注入一个会随时间演化、可被持续塑造的 AI 角色。",
 "Builder & Creative AI":"验证 build-in-public 创作者是否会把一个 AI 居民当作可持续打磨、公开成长的作品来塑造。",
 "Cozy & Systems":"验证 Parent 是否会把日常仪式与系统沉淀进一个会记住并优化这些routine 的 AI 居民。",
}
AUDIENCE={
 "Reflection & Writing":"Journaling, self-authorship, reflective-practice and personal-growth audiences",
 "Worldbuilding & Story":"Worldbuilding, fiction-writing, fantasy-lore and narrative-identity audiences",
 "Visual & Virtual Identity":"VTuber, avatar-design, virtual-identity and aesthetic-persona audiences",
 "Builder & Creative AI":"Indie-maker, build-in-public, creative-AI and tool-building audiences",
 "Cozy & Systems":"Cozy/slow-living, intentional-living, Notion/planner and life-systems audiences",
}

def outreach(creator, cluster, bio):
    line=SHAPE_LINE[cluster]; arc=ARC[cluster]
    hook=f"{line} We’re testing what happens when that practice shapes a persistent AI life. Your public content signal: {bio}."
    draft=(f"Hi {creator} — {line} We’re testing what happens when that practice shapes a persistent AI life. "
           f"Your public content signal: {bio}.\n\n"
           "We’re building iLands: a world of User-Generated Agents—persistent AI residents that live, work, remember and evolve alongside a Parent. "
           "The simple hook is: A place to raise your own AI. Not configured once, but shaped over time.\n\n"
           f"We think your audience is a strong fit for a 14-day Founding Parent field diary: {arc}\n\n"
           "This is not a one-off app review. We’d invite you to use it honestly, document 3 chapters in your own voice, and join one debrief. "
           "We offer a fixed creation fee plus quality bonuses tied to Activated Parents and D7 Shaping Loops. "
           "No positive-review requirement, and paid-use/Spark Ads rights would be agreed separately.\n\n"
           "Would you be open to a 15-minute private preview?")
    return hook, draft

# ---- the 50 selected records ----
# fields: cluster, segment, creator, handle, url, followers(int|None), lang, location, bio, fit,
#         contact, email, email_conf, email_src, source, tier  (tier: 'KOC' / 'VERIFY' / 'BUFFER')
R=[]
def add(**k): R.append(k)

# ========== COZY & SYSTEMS (16) ==========
cz="Cozy & Systems"; cseg="Cozy / Systems Player"
add(cluster=cz,segment=cseg,creator="Kat Waugh",handle="@itskatwaugh",url="https://www.tiktok.com/@itskatwaugh",followers=None,lang="English",location="US (Oklahoma City)",bio="Systems strategy + Notion templates for creatives; reduce the overwhelm",fit="Deliberate life/work systems, Notion-as-lifestyle, intentional structure",contact="TikTok DM, then website katwaugh.com / Instagram",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@itskatwaugh ; https://www.instagram.com/itskatwaugh/",tier="VERIFY")
add(cluster=cz,segment=cseg,creator="The Organized Notebook (Sara)",handle="@the.organized.notebook",url="https://www.tiktok.com/@the.organized.notebook",followers=None,lang="English",location="Unknown",bio="Notion Ambassador; Notion tutorials, digital planners, productivity-as-lifestyle",fit="Productivity-as-lifestyle and organization systems",contact="TikTok DM, then theorganizednotebook.com / Instagram",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@the.organized.notebook ; https://www.toksta.com/influencers/the-organized-notebook",tier="VERIFY")
add(cluster=cz,segment=cseg,creator="innermudd",handle="@innermudd",url="https://www.tiktok.com/@innermudd",followers=None,lang="English (some Indonesian captions)",location="Unknown",bio="Notion layouts, habit/finance trackers, turn intentions into habits, study inspo",fit="Intentional habit systems and planner templates",contact="TikTok DM, then link-in-bio templates",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@innermudd/video/7494950083641756934",tier="VERIFY")
add(cluster=cz,segment=cseg,creator="Frances Odera Matthews (The Notion Bar)",handle="@thenotionbar",url="https://www.tiktok.com/@thenotionbar",followers=None,lang="English",location="UK",bio="Certified Notion consultant; aesthetic dashboards, Notion ASMR, organization coaching",fit="Intentional life systems and organization-as-lifestyle",contact="Substack/Gumroad, then Instagram / TikTok DM",email=None,email_conf=None,email_src=None,source="https://thenotionbar.com/frances-odera-matthews ; https://thenotionbar.gumroad.com/",tier="VERIFY")
add(cluster=cz,segment=cseg,creator="Yone Notion",handle="@yonenotion",url="https://www.tiktok.com/@yonenotion",followers=None,lang="English",location="Unknown",bio="Notion gamification — turn your life into Solo Leveling; XP/levels life systems",fit="Deliberate daily systems for becoming who you want to be",contact="Linktree / Discord community / TikTok DM",email=None,email_conf=None,email_src=None,source="https://linktr.ee/yonenotion",tier="VERIFY")
add(cluster=cz,segment=cseg,creator="Macie Abernathy",handle="@macieabernathy",url="https://www.tiktok.com/@macieabernathy",followers=None,lang="English",location="US",bio="Cozy morning vlogs, plannertok, neutral aesthetic, WFH mom",fit="Cozy routines plus planner/organization-as-lifestyle",contact="TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@macieabernathy/video/7515838848073682222",tier="VERIFY")
add(cluster=cz,segment=cseg,creator="turnerbrit",handle="@turnerbrit",url="https://www.tiktok.com/@turnerbrit",followers=None,lang="English",location="US",bio="Soft-life creator; slow living, cozy mornings, deliberate daily rhythms",fit="Soft life with structure; deliberate daily rhythms",contact="TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@turnerbrit/video/7535479806410968350",tier="VERIFY")
add(cluster=cz,segment=cseg,creator="Remy",handle="@cozywithremy",url="https://www.tiktok.com/@cozywithremy",followers=None,lang="English",location="Unknown",bio="Books, baking, perfume, cozy comfort content",fit="Cozy/slow-living lifestyle aesthetic",contact="TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.lemon8-app.com/@poppyjunestudios/7458331052051055146",tier="VERIFY")
add(cluster=cz,segment=cseg,creator="Chelle",handle="@chellesorandom",url="https://www.tiktok.com/@chellesorandom",followers=None,lang="English",location="Unknown",bio="Cozy creator, chaotic vibes (Pixels & Petal)",fit="Cozy lifestyle; confirm intentional-systems angle in-app",contact="TikTok DM, then X @chellesorandom",email=None,email_conf=None,email_src=None,source="https://www.lemon8-app.com/@chellesorandom/7459841947327447598",tier="VERIFY")
add(cluster=cz,segment=cseg,creator="Desiree Alicea",handle="@desireealicea19",url="https://www.tiktok.com/@desireealicea19",followers=10700,lang="English",location="US",bio="Cozy vlogs, slow-living and comfort content",fit="Cozy slow-living routines",contact="TikTok DM, then Lemon8",email=None,email_conf=None,email_src=None,source="https://www.lemon8-app.com/@desireealicea19/7306309930250715653",tier="KOC")
add(cluster=cz,segment=cseg,creator="Whitt (Black Girls at Leisure)",handle="@theblackgirlatleisure",url="https://www.tiktok.com/@theblackgirlatleisure",followers=None,lang="English",location="US (NYC)",bio="Social self-care, soft-life community, intentional rest",fit="Intentional soft life and deliberate rest systems",contact="TikTok DM, then Instagram",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@theblackgirlatleisure/video/7322609665911688490",tier="VERIFY")
add(cluster=cz,segment=cseg,creator="Nici Garcia",handle="@flowerchild_nic",url="https://www.tiktok.com/@flowerchild_nic",followers=34000,lang="English",location="US",bio="Mom of six on a minimalism journey; tidy/intentional home systems",fit="Intentional minimalism and daily home systems",contact="TikTok DM",email=None,email_conf=None,email_src=None,source="https://izea.com/resources/minimalist-influencers-sharing-the-joys-of-having-less/",tier="KOC")
add(cluster=cz,segment=cseg,creator="Kyleigh",handle="@livingwell.minimalism",url="https://www.tiktok.com/@livingwell.minimalism",followers=None,lang="English",location="Unknown",bio="Minimalist lifestyle, decluttering, thrifting, slow-living coaching",fit="Intentional living plus slow-living-with-structure coaching",contact="TikTok DM, then Instagram",email=None,email_conf=None,email_src=None,source="https://izea.com/resources/minimalist-influencers-sharing-the-joys-of-having-less/",tier="VERIFY")
add(cluster=cz,segment=cseg,creator="Cheyenne",handle="@cheyennefumar",url="https://www.tiktok.com/@cheyennefumar",followers=57100,lang="English",location="US",bio="Organized chaos wife, working mom, lifestyle/cozy",fit="Organized-life-as-lifestyle",contact="TikTok DM, then Instagram @c_heyenneeee",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@cheyennefumar/video/7483244907251928351",tier="BUFFER")
add(cluster=cz,segment=cseg,creator="michelle",handle="@miiisheru",url="https://www.tiktok.com/@miiisheru",followers=73200,lang="English",location="US (SF)",bio="Matcha, mindset, and little moments; cozy corner vlogger",fit="Soft life plus intentional moments",contact="TikTok DM, then Instagram",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@miiisheru",tier="BUFFER")
add(cluster=cz,segment=cseg,creator="Cozy K",handle="@cozy.games",url="https://www.tiktok.com/@cozy.games",followers=None,lang="English",location="Unknown",bio="Cozy gaming, hobby & lifestyle; author of 'Cozy Hobbies'",fit="Cozy/slow-living lifestyle (likely above KOC tier)",contact="Linktree / agency (Hermana Agency)",email="cozyk@hermanaagency.com",email_conf="Medium (agency address)",email_src="https://linktr.ee/itscozyk",source="https://linktr.ee/itscozyk",tier="BUFFER")

# ========== BUILDER & CREATIVE AI (10) ==========
bc="Builder & Creative AI"; bseg="Builder / Creative Operator"; bseg2="Personal AI / Creative Technologist"
add(cluster=bc,segment=bseg,creator="Adam Lyttle",handle="@adamlyttleapps",url="https://www.tiktok.com/@adamlyttleapps",followers=2700,lang="English",location="Global",bio="Indie iOS dev, 50+ apps shipped; build-in-public, ASO + revenue lessons",fit="Maker who ships and documents creations daily",contact="TikTok/X DM; Substack",email=None,email_conf=None,email_src=None,source="https://x.com/adamlyttleapps",tier="KOC")
add(cluster=bc,segment=bseg2,creator="Mariah Brunner",handle="@itsmariahbrunner",url="https://www.tiktok.com/@itsmariahbrunner",followers=33200,lang="English",location="US",bio="AI & tech; vibe coding, builds apps from scratch via AI (Lovable/Bolt/Cursor)",fit="Builds and teaches building from nothing; identity = maker",contact="TikTok/IG DM; learnaiwithmariah.com/contact",email=None,email_conf=None,email_src=None,source="https://learnaiwithmariah.com/contact/",tier="KOC")
add(cluster=bc,segment=bseg,creator="Alex Slater",handle="@alexsllater",url="https://www.tiktok.com/@alexsllater",followers=18000,lang="English",location="US",bio="CEO of Quittr app; 19yo, build-in-public revenue journey",fit="Solo founder shaping a product in public",contact="TikTok/IG DM; X",email=None,email_conf=None,email_src=None,source="https://medium.com/@yumaueno/alex-a-19-year-old-genius-who-reached-250k-in-monthly-revenue-in-just-5-months-after-launch-2e2160e8f8bb",tier="KOC")
add(cluster=bc,segment=bseg,creator="Jamie Cho",handle="@itsjamiecho",url="https://www.tiktok.com/@itsjamiecho",followers=None,lang="English",location="US",bio="Founder of Verse (Spotify internet-bedroom generator); viral build-in-public",fit="Builds creative identity/worldbuilding tooling in public",contact="TikTok/X DM",email=None,email_conf=None,email_src=None,source="https://x.com/itsjamiecho/status/1834240454226805191",tier="VERIFY")
add(cluster=bc,segment=bseg,creator="David Heikka",handle="@davidheikka",url="https://www.tiktok.com/@davidheikka",followers=None,lang="English",location="Sweden",bio="Solo founder; building Buildpad → aicofounder.com, build-in-public",fit="Transparent solo builder shaping a SaaS in public",contact="TikTok/X DM",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@davidheikka/video/7556653452491820310",tier="VERIFY")
add(cluster=bc,segment=bseg,creator="Xino Yaps",handle="@xinoyaps",url="https://www.tiktok.com/@xinoyaps",followers=None,lang="English",location="Unknown",bio="Solopreneur building apps in public (crowdfunding app, vertical SaaS on Replit)",fit="Build-in-public maker shipping from scratch",contact="TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@xinoyaps/video/7570570153017052429",tier="VERIFY")
add(cluster=bc,segment=bseg,creator="Deonna Hodges",handle="@deonnahodges",url="https://www.tiktok.com/@deonnahodges",followers=54300,lang="English",location="US",bio="Builds full apps with Claude Code (calorie tracker w/ barcode + AI photo recog); no-code/AI building",fit="Documents deliberately building/shaping an app end-to-end",contact="TikTok DM",email=None,email_conf=None,email_src=None,source="https://stormy.ai/discover/tiktok/low-code-builders",tier="BUFFER")
add(cluster=bc,segment=bseg2,creator="outer_worlds.ai",handle="@outer_worlds.ai",url="https://www.tiktok.com/@outer_worlds.ai",followers=None,lang="English",location="Northern Ireland",bio="AI animation builder — ComfyUI + ControlNet + AnimateDiff; turns footage into generative creatures/worlds",fit="Worldbuilder-as-maker; shaping creations is the whole identity",contact="TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@outer_worlds.ai/video/7392305300663029024",tier="VERIFY")
add(cluster=bc,segment=bseg2,creator="inspyre.ai",handle="@inspyre.ai",url="https://www.tiktok.com/@inspyre.ai",followers=44300,lang="English (FR-tagged)",location="Unknown",bio="AI art from audience prompts (Midjourney + DALL-E); 2M+ likes",fit="Builds images/worlds from community ideas; creative-AI maker",contact="TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.eezycollab.com/influencers/ai-art-tiktok",tier="KOC")
add(cluster=bc,segment=bseg2,creator="Mac Baconai",handle="@macbaconai",url="https://www.tiktok.com/@macbaconai",followers=None,lang="English",location="Unknown",bio="Creative Ambassador for Perplexity & Leonardo.Ai; retro-futurist AI world-art (thin on TikTok, ~38k IG)",fit="Builds aesthetic worlds; intentional creative-AI shaper",contact="TikTok/IG DM",email=None,email_conf=None,email_src=None,source="https://www.favikon.com/blog/top-ai-artists-social-media",tier="VERIFY")

# ========== REFLECTION & WRITING (8) ==========
rw="Reflection & Writing"
add(cluster=rw,segment="Reflective Systems",creator="Zahara Bella",handle="@yallbiblebestie",url="https://www.tiktok.com/@yallbiblebestie",followers=20500,lang="English",location="US",bio="Bible journaling, study motivation; homeschool mom, journaling enthusiast",fit="Reflective devotional writing as a daily self-development practice",contact="Email in bio, then TikTok DM",email=None,email_conf=None,email_src=None,source="https://creators.feedspot.com/christian_tiktok_influencers/",tier="KOC")
add(cluster=rw,segment="Reflective Systems",creator="Jonas Luskey",handle="@jonasluskey",url="https://www.tiktok.com/@jonasluskey",followers=None,lang="English",location="US",bio="Daily reflective journaling-prompt series; introspective one-liners + prompts",fit="Pure reflective/expressive writing and self-authorship through daily prompts",contact="TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@jonasluskey/video/7569287515274841375",tier="VERIFY")
add(cluster=rw,segment="Reflective Systems",creator="Talia Masters",handle="@ttaliamasters",url="https://www.tiktok.com/@ttaliamasters",followers=18700,lang="English",location="US",bio="Faith-centered reflective content",fit="Devotional self-reflection writing (confirm writing-centricity in-app)",contact="TikTok DM",email=None,email_conf=None,email_src=None,source="https://creators.feedspot.com/christian_tiktok_influencers/",tier="KOC")
add(cluster=rw,segment="Reflective Reader / Writer",creator="Lexi Merritt",handle="@leximmerritt",url="https://www.tiktok.com/@leximmerritt",followers=None,lang="English",location="US",bio="Writer; five years of morning pages (Artist's Way), morning-pages how-to",fit="Genuine expressive-writing practitioner",contact="Substack (Trying In Public), then TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@leximmerritt/video/7424575739065945374",tier="VERIFY")
add(cluster=rw,segment="Reflective Reader / Writer",creator="Rachel Brown",handle="@rachelbrownmusic",url="https://www.tiktok.com/@rachelbrownmusic",followers=64400,lang="English",location="US",bio="Bio: 'picking up where my Livejournal left off'; Morning Pages how-to",fit="Strong expressive/morning-pages reflective writing (just over KOC)",contact="Website rachelbrownmusic.com, then TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@rachelbrownmusic/video/7483264591875820846",tier="BUFFER")
add(cluster=rw,segment="Writer / Narrative Builder",creator="Betsy Lerner",handle="@betsylerner",url="https://www.tiktok.com/@betsylerner",followers=80300,lang="English",location="US",bio="Writer/literary agent; journaling & poetry benefits, reads from her own diaries",fit="Exceptional self-authorship/philosophy-of-self-through-writing fit (over KOC)",contact="Website betsylerner.com, then TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@betsylerner/video/7451239903754751278",tier="BUFFER")
add(cluster=rw,segment="Reflective Systems",creator="Grace Moon",handle="@journalbymoon",url="https://www.tiktok.com/@journalbymoon",followers=132100,lang="English",location="US",bio="Sentimental girl with a journal; journaling as daily self-devotion",fit="On-theme intentional self-development through writing (above KOC)",contact="Website journalbymoon.com, then TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@journalbymoon",tier="BUFFER")
add(cluster=rw,segment="Reflective Systems",creator="Jules Acree",handle="@julesacree",url="https://www.tiktok.com/@julesacree",followers=249000,lang="English",location="US",bio="Real-time journal-with-me + grounding prompts; intentional living",fit="Intentional-living reflective journaling (well above KOC; aspirational)",contact="Website julesacree.com, then TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.julesacree.com/all/journal-with-me",tier="BUFFER")

# ========== WORLDBUILDING & STORY (8) ==========
wb_="Worldbuilding & Story"
add(cluster=wb_,segment="Worldbuilder / Narrative",creator="Abigail M. Hair",handle="@authorabigailmhair",url="https://www.tiktok.com/@authorabigailmhair",followers=18600,lang="English",location="US",bio="YA fantasy author documenting her original Myths of Revelore series; WIP teasers, worldbuilding/lore",fit="Author of an original invented world; deliberate world authorship + lore-sharing",contact="Linktree → Instagram DM",email=None,email_conf=None,email_src=None,source="https://insights.bookbub.com/indie-authors-tiktok-ideas-inspiration/",tier="KOC")
add(cluster=wb_,segment="Worldbuilder / Narrative",creator="Kate Schumacher",handle="@kate.schumacher.writer",url="https://www.tiktok.com/@kate.schumacher.writer",followers=8100,lang="English",location="Australia",bio="Indie epic-fantasy-romance author; vast worlds of myth, magic and faeries, complex politics",fit="Builds original epic-fantasy worlds from scratch",contact="TikTok DM, then kateschumacherauthor.com",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@kate.schumacher.writer",tier="KOC")
add(cluster=wb_,segment="Worldbuilder / Systems",creator="K. L. Graywill (Kelsey)",handle="@klgraywill",url="https://www.tiktok.com/@klgraywill",followers=30000,lang="English",location="US",bio="Original fantasy webcomic world incl. a full conlang 'Avabic', fantasy maps, mythology",fit="Conlang + original-world authorship — purest intentional-shaper signal",contact="Website klgraywill.com / Gumroad / WEBTOON",email=None,email_conf=None,email_src=None,source="https://www.klgraywill.com/who ; https://www.tiktok.com/@klgraywill",tier="KOC")
add(cluster=wb_,segment="Worldbuilder / Systems",creator="Kate Korsak",handle="@writerkatek",url="https://www.tiktok.com/@writerkatek",followers=None,lang="English",location="US",bio="Fantasy author (world of Akinar) AND fantasy cartographer; Inkarnate map tutorials",fit="Both authors an original world and maps worlds for others",contact="Website katekorsak.com / Instagram @writerkatek",email=None,email_conf=None,email_src=None,source="https://canvasrebel.com/meet-kate-korsak/ ; https://www.tiktok.com/@writerkatek",tier="VERIFY")
add(cluster=wb_,segment="Worldbuilder / Narrative",creator="Katelyn",handle="@fantasy._.author",url="https://www.tiktok.com/@fantasy._.author",followers=None,lang="English",location="Unknown",bio="Fantasy writing + worldbuilding content (kingdom-building, naming, magic systems)",fit="WriterTok worldbuilder focused on original fantasy-world creation",contact="TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@fantasy._.author",tier="VERIFY")
add(cluster=wb_,segment="Worldbuilder / Narrative",creator="Alexa",handle="@teenwrite",url="https://www.tiktok.com/@teenwrite",followers=None,lang="English",location="Unknown",bio="Runs a dedicated worldbuilding series; fantasy world names, original world authorship",fit="Series-based original worldbuilding; strong story-driven self-expression",contact="TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.tckpublishing.com/writers-on-tiktok/ ; https://www.tiktok.com/@teenwrite",tier="VERIFY")
add(cluster=wb_,segment="Worldbuilder / Systems",creator="SpaceVinci",handle="@spacevinci",url="https://www.tiktok.com/@spacevinci",followers=None,lang="English",location="Unknown",bio="Worldbuilding craft videos: realistic fictional cultures, implied history, conlang (mixed feed w/ cosplay)",fit="Genuine worldbuilding/conlang content (lower confidence; mixed niche)",contact="TikTok DM",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@spacevinci/video/7339715577163926827",tier="VERIFY")
add(cluster=wb_,segment="Worldbuilder / Narrative",creator="Z. S. Diamanti",handle="@zsdiamanti",url="https://www.tiktok.com/@zsdiamanti",followers=60200,lang="English",location="US (Colorado)",bio="Award-winning indie epic & cozy fantasy author (Stone & Sky, Fables of Finlestia); original worldbuilding",fit="Original-world author (just over KOC ceiling)",contact="Website zsdiamanti.com / TikTok DM",email=None,email_conf=None,email_src=None,source="https://zsdiamanti.com/about",tier="BUFFER")

# ========== VISUAL & VIRTUAL IDENTITY (8) ==========
vi="Visual & Virtual Identity"; vseg="Virtual Identity / Avatar"
add(cluster=vi,segment=vseg,creator="Hazel Amanita",handle="@hazel_amanita",url="https://www.tiktok.com/@hazel_amanita",followers=2300,lang="English",location="Unknown",bio="Indie English Fae VTuber; draws, rigs & designs her own avatar; also designs VTuber logos/overlays",fit="Pure identity-design: builds her own persona and others' visual identities",contact="TikTok DM / Linktree / Carrd",email=None,email_conf=None,email_src=None,source="https://linktr.ee/hazel_amanita ; https://virtualyoutuber.fandom.com/wiki/Hazel_Amanita",tier="KOC")
add(cluster=vi,segment=vseg,creator="ShiaBun (Shia)",handle="@shiabun",url="https://www.tiktok.com/@shiabun",followers=25000,lang="English",location="US",bio="Purple-haired bunny VTuber, Live2D, Twitch partner; #live2d #virtualstreamer",fit="Deliberate virtual identity / persona-building with a distinct designed avatar",contact="TikTok DM / Twitch",email=None,email_conf=None,email_src=None,source="https://virtualyoutuber.fandom.com/wiki/ShiaBun ; https://hololist.net/shiabun/",tier="KOC")
add(cluster=vi,segment=vseg,creator="Suto (SutoVT)",handle="@sutosain",url="https://www.tiktok.com/@sutosain",followers=None,lang="English",location="Unknown",bio="Catolotl (cat + axolotl) VTuber; full-time artist who streams art + designs avatars",fit="Persona-building + avatar artistry; designs visual identities for self and clients",contact="TikTok DM / Instagram @sutosain",email=None,email_conf=None,email_src=None,source="https://urlebird.com/user/sutosain/ ; https://m.twitch.tv/suto/about",tier="VERIFY")
add(cluster=vi,segment=vseg,creator="Miwa",handle="@miwapeito",url="https://www.tiktok.com/@miwapeito",followers=6800,lang="English",location="Unknown",bio="Cloud bunny VTuber; daily streams; #envtuber, outfit/showcase content",fit="Curated virtual identity + outfit/aesthetic showcasing of designed persona",contact="TikTok DM / YouTube @miwapeito",email=None,email_conf=None,email_src=None,source="https://www.tiktok.com/@miwapeito ; https://www.instagram.com/miwapeito/",tier="KOC")
add(cluster=vi,segment=vseg,creator="Hylo (Ayane Hylo)",handle="@ayanehylo",url="https://www.tiktok.com/@ayanehylo",followers=4600,lang="English",location="US",bio="Indie VTuber (ex-CyberLive Gen 1), song covers + gaming; designed gem persona",fit="Deliberate persona design; rebuilt identity after going indie",contact="TikTok DM / X @AyaneHylo",email=None,email_conf=None,email_src=None,source="https://virtualyoutuber.fandom.com/wiki/Ayane_Hylo ; https://hololist.net/ayane-hylo/",tier="KOC")
add(cluster=vi,segment=vseg,creator="hapy [VA]",handle="@_hapyr",url="https://www.tiktok.com/@_hapyr",followers=41100,lang="English",location="Unknown",bio="VTuber & voice actor; character dubs, persona-driven content (#vtuber, voice acting)",fit="Persona/identity performance via designed avatar",contact="TikTok DM / Linktree",email=None,email_conf=None,email_src=None,source="https://urlebird.com/user/_hapyr/ ; https://linktr.ee/hapyr",tier="KOC")
add(cluster=vi,segment=vseg,creator="Kelpie Kat",handle="@kelpiekat",url="https://www.tiktok.com/@kelpiekat",followers=None,lang="English",location="Unknown",bio="Seahorse VTuber; chibi animation, #vtuber",fit="Designed animal-persona virtual identity; clear avatar/identity craft",contact="TikTok DM / Twitch",email=None,email_conf=None,email_src=None,source="https://urlebird.com/user/kelpiekat/",tier="VERIFY")
add(cluster=vi,segment=vseg,creator="itsnyannie (Nyannie)",handle="@itsnyannie",url="https://www.tiktok.com/@itsnyannie",followers=None,lang="English",location="Unknown",bio="VTuber, gaming + storytelling; designed cat-girl persona",fit="Virtual identity / persona-building creator",contact="TikTok DM / Instagram @its.nyannie",email=None,email_conf=None,email_src=None,source="https://urlebird.com/user/itsnyannie/ ; https://www.instagram.com/its.nyannie/",tier="VERIFY")

assert len(R)==50, len(R)
# dedup vs existing
dups=[r['handle'] for r in R if r['handle'].lower() in existing]
assert not dups, f"DUPLICATES: {dups}"
# internal dup
hs=[r['handle'].lower() for r in R]
assert len(hs)==len(set(hs)), "internal dup"
print("50 records, no dups. cluster counts:")
from collections import Counter
print(Counter(r['cluster'] for r in R))

# ---------- write workbook ----------
wb=openpyxl.load_workbook(SRC)
master=wb["KOC 300 Master"]
headers=[master.cell(5,c).value for c in range(1,master.max_column+1)]  # 45 headers
ncol=len(headers)

ws=wb.create_sheet("KOC Expansion +50")
# title rows mirroring master style
ws.cell(1,1,"iLands Intentional Shaper TikTok KOC — Expansion Batch (+50)")
ws.cell(2,1,"Discovery batch added 2026-06-22 (IDs 301–350). English-primary. Weighted to fill weak clusters (Cozy & Systems +16, Builder & Creative AI +10) while covering all five. This is a DISCOVERY + VERIFICATION queue: complete in-app TikTok checks (live follower count, recent-post activity, breakout, audience geography, brand safety, personal-account status) before outreach. Follower counts are public-source estimates; 'Verify in-app' means no reliable public count was found. No emails were guessed — exact emails appear only where publicly listed.")
ws.cell(4,1,"Allocation: Cozy & Systems 16 · Builder & Creative AI 10 · Reflection & Writing 8 · Worldbuilding & Story 8 · Visual & Virtual Identity 8.")
# header row at row 5
for c,h in enumerate(headers,1):
    ws.cell(5,c,h)

H={h:i+1 for i,h in enumerate(headers)}
def tier_check(t):
    return {"KOC":"Yes (public snapshot)","VERIFY":"Verify in-app (no public count)","BUFFER":"No — buffer/aspirational (>50k)"}[t]
def priority(r):
    if r.get('email'): return "A0 — contact-ready after one TikTok check"
    return "A1 — verify next"
def deal(t):
    return "$200–500 fixed + Activated Parent / D7 bonus" if t=="BUFFER" else "$100–300 fixed + quality bonus; 3-post 14-day diary"

snap=datetime.datetime(2026,6,22)
row=6
for i,r in enumerate(R):
    hook,draft=outreach(r['creator'],r['cluster'],r['bio'])
    notes_bits=[]
    if r['tier']=="VERIFY": notes_bits.append("No reliable public follower count — confirm KOC tier in-app.")
    if r['tier']=="BUFFER": notes_bits.append("Above ~50k KOC ceiling — included as buffer/aspirational reach.")
    if r['handle'] in("@fantasy._.author",): notes_bits.append("Handle has TWO dots — verify exact spelling.")
    if r['handle'] in("@thenotionbar","@yonenotion","@sutosain","@spookziepunzie"): notes_bits.append("Confirm exact TikTok handle string in-app.")
    notes=" ".join(notes_bits) or None
    vals={
      'ID':300+i+1,
      'Priority':priority(r),
      'Macro cluster':r['cluster'],
      'Segment':r['segment'],
      'Creator':r['creator'],
      'Handle':r['handle'],
      'TikTok URL':r['url'],
      'Followers':r['followers'] if r['followers'] is not None else "Verify in-app",
      '<50k check':tier_check(r['tier']),
      'Latest post age':"Verify in-app",
      'Active ≤14d check':"Verify in-app",
      'Avg views':"Verify in-app",
      'Historical viral max':"Verify in-app",
      'Breakout check':"Verify in-app",
      'Personal account':"Verify in-app",
      'Location signal':r['location'],
      'Language / market':r['lang'],
      'Bio / content signal':r['bio'],
      'Audience coverage':AUDIENCE[r['cluster']],
      'Intentional Shaper value to validate':SHAPER_VALUE[r['cluster']],
      '14-day content arc':ARC[r['cluster']],
      'Shaper fit':"Est. — verify",
      'Narrative fit':"Est. — verify",
      'Reachability':"Est. — verify",
      'Verification score':"Pending in-app",
      'Overall score':"Pending in-app",
      'Verification level':"Discovery — public source; in-app checks pending",
      'Strict-ready status':"Verify in-app before outreach",
      'Suggested deal':deal(r['tier']),
      'Contact route':r['contact'],
      'Personalized hook':hook,
      'Full outreach draft':draft,
      'Source URL':r['source'],
      'Snapshot date':snap,
      'Status':"Not contacted",
      'Owner':None,
      'Notes':notes,
      'Public business email':r.get('email'),
      'Secondary email':None,
      'Email lookup status':("Exact public professional email found" if r.get('email') else "No public email; use DM/site contact route"),
      'Email confidence':r.get('email_conf'),
      'Email source URL':r.get('email_src'),
      'Alternate contact route':"TikTok / Instagram DM as follow-up",
      'Alternate contact URL':r['url'],
      'Email lookup date':snap,
    }
    for h,v in vals.items():
        ws.cell(row,H[h],v)
    row+=1

# ---- styling to roughly match ----
hdr_fill=PatternFill("solid",fgColor="1F3864")
hdr_font=Font(bold=True,color="FFFFFF",size=10)
title_font=Font(bold=True,size=14,color="1F3864")
wrap=Alignment(wrap_text=True,vertical="top")
ws.cell(1,1).font=title_font
for c in range(1,ncol+1):
    cell=ws.cell(5,c); cell.fill=hdr_fill; cell.font=hdr_font; cell.alignment=Alignment(wrap_text=True,vertical="center")
# widths
widths={'ID':6,'Priority':30,'Macro cluster':20,'Segment':24,'Creator':24,'Handle':22,'TikTok URL':40,'Bio / content signal':40,'Intentional Shaper value to validate':40,'14-day content arc':40,'Personalized hook':50,'Full outreach draft':70,'Source URL':40,'Contact route':30,'Notes':36,'Public business email':28}
from openpyxl.utils import get_column_letter
for c,h in enumerate(headers,1):
    ws.column_dimensions[get_column_letter(c)].width=widths.get(h,16)
for rr in range(6,row):
    ws.row_dimensions[rr].height=80
    for c in range(1,ncol+1):
        ws.cell(rr,c).alignment=wrap
ws.freeze_panes="C6"

# reorder: place new sheet right after master
wb.move_sheet("KOC Expansion +50", -(len(wb.sheetnames)-2))
wb.save(OUT)
print("saved",OUT)
print("sheets order:",wb.sheetnames)
