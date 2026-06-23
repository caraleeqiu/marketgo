import openpyxl, warnings, json, datetime, re
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
warnings.filterwarnings("ignore")

BASE="/home/user/marketgo/iLands_Intentional_Shaper_TikTok_KOC_350_with_Emails_2026-06-22.xlsx"
OUT="/home/user/marketgo/iLands_Intentional_Shaper_TikTok_KOC_with_Emails_2026-06-22.xlsx"
existing=set(json.load(open("/home/user/marketgo/.research/all_handles_350.json")))

rows=[]
for ln in open("/home/user/marketgo/.research/verified_emails.txt"):
    ln=ln.strip()
    if not ln or ln.startswith("#"): continue
    p=[x.strip() for x in ln.split("|")]
    if len(p)<8: continue
    cluster,creator,handle,url,foll,bio,email,src=p[:8]
    rows.append(dict(cluster=cluster,creator=creator,handle=handle,url=url,foll=foll,bio=bio,email=email,src=src))

# dedup
dups=[r['handle'] for r in rows if r['handle'].lower() in existing]
print("dups vs existing 350:",dups)
seen=set(); uniq=[]
for r in rows:
    h=r['handle'].lower()
    if h in seen or h in existing: continue
    seen.add(h); uniq.append(r)
rows=uniq
print("unique verified-email records:",len(rows))

SHAPE_LINE={
 "Reflection & Writing":"You already turn inner life into a practice.",
 "Worldbuilding & Story":"You already author whole worlds and make them feel inhabited.",
 "Visual & Virtual Identity":"You already design and inhabit a deliberate visual identity.",
 "Builder & Creative AI":"You already build things in public and shape them over time.",
 "Cozy & Systems":"You already turn intentional routines into living systems.",
}
ARC={
 "Reflection & Writing":"14-night shaping diary: one 20-minute evening reflection becomes a visible choice, memory or relationship change in the Agent.",
 "Worldbuilding & Story":"14-day worldbuilding diary: each session adds canon, characters and history that your AI resident remembers and acts on.",
 "Visual & Virtual Identity":"14-day identity diary: each session shapes how your AI resident looks, presents and evolves its persona.",
 "Builder & Creative AI":"14-day build-in-public diary: each session ships a visible change to an AI resident you are raising from scratch.",
 "Cozy & Systems":"14-day cozy-systems diary: one calm evening ritual becomes a routine your AI resident adopts, remembers and refines.",
}
SHAPER_VALUE={
 "Reflection & Writing":"验证 iLands 是否能成为一种会复利的数字关系：Parent 把判断、人生语境与反思沉淀进一个持续成长的 AI。",
 "Worldbuilding & Story":"验证 Parent 是否会把世界观、设定与叙事记忆投入一个持续演化、能记住 canon 的 AI 居民。",
 "Visual & Virtual Identity":"验证 Parent 是否愿意把虚拟形象与身份审美注入一个会随时间演化、可被持续塑造的 AI 角色。",
 "Builder & Creative AI":"验证 build-in-public 创作者是否会把一个 AI 居民当作可持续打磨、公开成长的作品来塑造。",
 "Cozy & Systems":"验证 Parent 是否会把日常仪式与系统沉淀进一个会记住并优化这些 routine 的 AI 居民。",
}
AUDIENCE={
 "Reflection & Writing":"Journaling, self-authorship, reflective-practice and personal-growth audiences",
 "Worldbuilding & Story":"Worldbuilding, fiction-writing, fantasy-lore and narrative-identity audiences",
 "Visual & Virtual Identity":"VTuber, avatar-design, virtual-identity and aesthetic-persona audiences",
 "Builder & Creative AI":"Indie-maker, build-in-public, creative-AI and tool-building audiences",
 "Cozy & Systems":"Cozy/slow-living, intentional-living, Notion/planner and life-systems audiences",
}
SEG={
 "Reflection & Writing":"Reflective Reader / Writer",
 "Worldbuilding & Story":"Worldbuilder / Narrative",
 "Visual & Virtual Identity":"Virtual Identity / Avatar",
 "Builder & Creative AI":"Personal AI / Creative Technologist",
 "Cozy & Systems":"Cozy / Systems Player",
}
def outreach(creator,cluster,bio):
    line=SHAPE_LINE[cluster]; arc=ARC[cluster]
    hook=f"{line} We’re testing what happens when that practice shapes a persistent AI life. Your public content signal: {bio}."
    draft=(f"Hi {creator} — {line} We’re testing what happens when that practice shapes a persistent AI life. "
           f"Your public content signal: {bio}.\n\n"
           "We’re building iLands: a world of User-Generated Agents—persistent AI residents that live, work, remember and evolve alongside a Parent. "
           "The simple hook is: A place to raise your own AI. Not configured once, but shaped over time.\n\n"
           f"We think your audience is a strong fit for a 14-day Founding Parent field diary: {arc}\n\n"
           "This is not a one-off app review. We’d invite you to use it honestly, document 3 chapters in your own voice, and join one debrief. "
           "We offer a fixed creation fee plus quality bonuses tied to Activated Parents and D7 Shaping Loops. "
           "No positive-review requirement, and paid-use/Spark Ads rights would be agreed separately.\n\n"
           "Would you be open to a 15-minute private preview?")
    return hook,draft

wb=openpyxl.load_workbook(BASE)
headers=[wb["KOC 300 Master"].cell(5,c).value for c in range(1,wb["KOC 300 Master"].max_column+1)]
ncol=len(headers); H={h:i+1 for i,h in enumerate(headers)}
ws=wb.create_sheet("KOC Email-Verified")
ws.cell(1,1,"iLands Intentional Shaper — Email-Verified Contacts (batch in progress)")
ws.cell(2,1,"Creators with a PUBLIC business email confirmed by viewing the source page directly (no guessed/constructed emails). IDs 351+. This batch is being expanded toward ~100; in-app TikTok checks (live followers, recency, brand safety) still pending before outreach. Most are indie authors / writers / journaling creators whose Linktree or site lists a contact email as text.")
ws.cell(4,1,"Note: many TikTok creators expose only a DM or a mailto button (no readable address); those are excluded here by design. Handles marked 'verify' need an exact in-app handle confirmation.")
for c,h in enumerate(headers,1): ws.cell(5,c,h)

snap=datetime.datetime(2026,6,22)
r0=6
for i,r in enumerate(rows):
    cl=r['cluster']
    if cl not in SHAPE_LINE: cl="Reflection & Writing"
    hook,draft=outreach(r['creator'],cl,r['bio'])
    foll=r['foll']
    foll_val=int(foll) if str(foll).isdigit() else (foll if foll and foll.lower()!="verify" else "Verify in-app")
    vals={
      'ID':350+i+1,'Priority':"A0 — contact-ready after one TikTok check",
      'Macro cluster':cl,'Segment':SEG[cl],'Creator':r['creator'],'Handle':r['handle'],'TikTok URL':r['url'],
      'Followers':foll_val,'<50k check':"Verify in-app",'Latest post age':"Verify in-app",'Active ≤14d check':"Verify in-app",
      'Avg views':"Verify in-app",'Historical viral max':"Verify in-app",'Breakout check':"Verify in-app",'Personal account':"Verify in-app",
      'Location signal':"Unknown",'Language / market':"English",'Bio / content signal':r['bio'],
      'Audience coverage':AUDIENCE[cl],'Intentional Shaper value to validate':SHAPER_VALUE[cl],'14-day content arc':ARC[cl],
      'Shaper fit':"Est. — verify",'Narrative fit':"Est. — verify",'Reachability':"High (public email)",
      'Verification score':"Pending in-app",'Overall score':"Pending in-app",
      'Verification level':"Discovery — public email confirmed on-page; in-app checks pending",
      'Strict-ready status':"Verify in-app before outreach",
      'Suggested deal':"$100–300 fixed + quality bonus; 3-post 14-day diary",
      'Contact route':"Public business email, then TikTok/Instagram DM",
      'Personalized hook':hook,'Full outreach draft':draft,'Source URL':r['src'],'Snapshot date':snap,
      'Status':"Not contacted",'Owner':None,
      'Notes':"Email confirmed by viewing source page directly. Confirm exact TikTok handle in-app." if "verify" in r['handle'].lower() else "Email confirmed by viewing source page directly.",
      'Public business email':r['email'],'Secondary email':None,
      'Email lookup status':"Exact public professional email found",'Email confidence':"High",
      'Email source URL':r['src'],'Alternate contact route':"TikTok / Instagram DM as follow-up",
      'Alternate contact URL':r['url'],'Email lookup date':snap,
    }
    for h,v in vals.items(): ws.cell(r0,H[h],v)
    r0+=1

# style
hdr_fill=PatternFill("solid",fgColor="2E5D34"); hdr_font=Font(bold=True,color="FFFFFF",size=10)
ws.cell(1,1).font=Font(bold=True,size=14,color="2E5D34")
for c in range(1,ncol+1):
    cell=ws.cell(5,c); cell.fill=hdr_fill; cell.font=hdr_font; cell.alignment=Alignment(wrap_text=True,vertical="center")
widths={'ID':6,'Priority':30,'Macro cluster':20,'Segment':24,'Creator':26,'Handle':24,'TikTok URL':40,'Bio / content signal':36,'Intentional Shaper value to validate':40,'14-day content arc':40,'Personalized hook':50,'Full outreach draft':70,'Source URL':40,'Contact route':30,'Notes':34,'Public business email':30,'Email source URL':36}
for c,h in enumerate(headers,1): ws.column_dimensions[get_column_letter(c)].width=widths.get(h,16)
wrap=Alignment(wrap_text=True,vertical="top")
for rr in range(6,r0):
    ws.row_dimensions[rr].height=80
    for c in range(1,ncol+1): ws.cell(rr,c).alignment=wrap
ws.freeze_panes="C6"
# place right after the +50 sheet
names=wb.sheetnames
idx_target=names.index("KOC Expansion +50")+1
wb.move_sheet("KOC Email-Verified", -(len(names)-1-idx_target))
wb.save(OUT)
print("saved",OUT)
print("order:",wb.sheetnames)
